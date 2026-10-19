from __future__ import annotations

import csv
import sqlite3
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "DBD_PCOM7E_Table.xlsx"


def clean_header(value: object) -> str:
    return str(value or "").replace("\u00a0", " ").strip()


def text(value: object) -> str:
    return str(value or "").strip()


def iso_date(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    return text(value)


def write_csv(filename: str, headers: list[str], rows: list[dict[str, object]]) -> None:
    with (ROOT / filename).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


workbook = load_workbook(SOURCE, data_only=True, read_only=True)
sheet = workbook.active
rows = list(sheet.iter_rows(values_only=True))
headers = [clean_header(value) for value in rows[0]]

expected = [
    "Student Number",
    "Student Name",
    "Exam Score",
    "Support",
    "Date of Birth",
    "Course Name",
    "Exam Boards",
    "Teacher Name",
]
if headers != expected:
    raise ValueError(f"Unexpected source headers: {headers!r}")

students: OrderedDict[int, dict[str, object]] = OrderedDict()
one_nf: list[dict[str, object]] = []
current_student: dict[str, object] | None = None

for source_row in rows[1:]:
    student_number, student_name, score, support, dob, course, board, teacher = source_row
    if student_number is not None:
        current_student = {
            "StudentNumber": int(student_number),
            "StudentName": text(student_name),
            "ExamScore": int(score),
            "Support": text(support),
            "DateOfBirth": iso_date(dob),
        }
        students[current_student["StudentNumber"]] = current_student.copy()
    if current_student is None:
        raise ValueError("A course row appeared before its student record.")
    one_nf.append(
        {
            **current_student,
            "CourseName": text(course),
            "ExamBoard": text(board),
            "TeacherName": text(teacher),
        }
    )

student_rows = list(students.values())
course_teacher: OrderedDict[str, str] = OrderedDict()
for row in one_nf:
    course_name = str(row["CourseName"])
    teacher_name = str(row["TeacherName"])
    previous = course_teacher.setdefault(course_name, teacher_name)
    if previous != teacher_name:
        raise ValueError(f"Course {course_name!r} has inconsistent teachers.")

two_nf_courses = [
    {"CourseName": course_name, "TeacherName": teacher_name}
    for course_name, teacher_name in course_teacher.items()
]
two_nf_student_courses = [
    {
        "StudentNumber": row["StudentNumber"],
        "CourseName": row["CourseName"],
        "ExamBoard": row["ExamBoard"],
    }
    for row in one_nf
]

teachers: OrderedDict[str, str] = OrderedDict()
boards: OrderedDict[str, str] = OrderedDict()
courses: OrderedDict[str, dict[str, str]] = OrderedDict()

for teacher_name in course_teacher.values():
    teachers.setdefault(teacher_name, f"T{len(teachers) + 1:02d}")
for row in one_nf:
    board_name = str(row["ExamBoard"])
    boards.setdefault(board_name, f"B{len(boards) + 1:02d}")
for course_name, teacher_name in course_teacher.items():
    courses[course_name] = {
        "CourseID": f"C{len(courses) + 1:02d}",
        "CourseName": course_name,
        "TeacherID": teachers[teacher_name],
    }

teacher_rows = [
    {"TeacherID": teacher_id, "TeacherName": teacher_name}
    for teacher_name, teacher_id in teachers.items()
]
board_rows = [
    {"ExamBoardID": board_id, "ExamBoardName": board_name}
    for board_name, board_id in boards.items()
]
course_rows = list(courses.values())
enrolment_rows = [
    {
        "StudentNumber": row["StudentNumber"],
        "CourseID": courses[str(row["CourseName"])]["CourseID"],
        "ExamBoardID": boards[str(row["ExamBoard"])],
    }
    for row in one_nf
]

write_csv(
    "1nf_flat.csv",
    [
        "StudentNumber",
        "StudentName",
        "ExamScore",
        "Support",
        "DateOfBirth",
        "CourseName",
        "ExamBoard",
        "TeacherName",
    ],
    one_nf,
)
write_csv("2nf_students.csv", list(student_rows[0]), student_rows)
write_csv("2nf_courses.csv", ["CourseName", "TeacherName"], two_nf_courses)
write_csv(
    "2nf_student_courses.csv",
    ["StudentNumber", "CourseName", "ExamBoard"],
    two_nf_student_courses,
)
write_csv("3nf_students.csv", list(student_rows[0]), student_rows)
write_csv("3nf_teachers.csv", ["TeacherID", "TeacherName"], teacher_rows)
write_csv("3nf_exam_boards.csv", ["ExamBoardID", "ExamBoardName"], board_rows)
write_csv("3nf_courses.csv", ["CourseID", "CourseName", "TeacherID"], course_rows)
write_csv(
    "3nf_enrolments.csv",
    ["StudentNumber", "CourseID", "ExamBoardID"],
    enrolment_rows,
)

schema_sql = """PRAGMA foreign_keys = ON;

CREATE TABLE Student (
    StudentNumber INTEGER PRIMARY KEY,
    StudentName TEXT NOT NULL,
    ExamScore INTEGER NOT NULL CHECK (ExamScore BETWEEN 0 AND 100),
    Support TEXT NOT NULL CHECK (Support IN ('Yes', 'No')),
    DateOfBirth TEXT NOT NULL
);

CREATE TABLE Teacher (
    TeacherID TEXT PRIMARY KEY,
    TeacherName TEXT NOT NULL UNIQUE
);

CREATE TABLE ExamBoard (
    ExamBoardID TEXT PRIMARY KEY,
    ExamBoardName TEXT NOT NULL UNIQUE
);

CREATE TABLE Course (
    CourseID TEXT PRIMARY KEY,
    CourseName TEXT NOT NULL UNIQUE,
    TeacherID TEXT NOT NULL,
    FOREIGN KEY (TeacherID) REFERENCES Teacher (TeacherID)
        ON UPDATE CASCADE ON DELETE RESTRICT
);

CREATE TABLE Enrolment (
    StudentNumber INTEGER NOT NULL,
    CourseID TEXT NOT NULL,
    ExamBoardID TEXT NOT NULL,
    PRIMARY KEY (StudentNumber, CourseID),
    FOREIGN KEY (StudentNumber) REFERENCES Student (StudentNumber)
        ON UPDATE CASCADE ON DELETE RESTRICT,
    FOREIGN KEY (CourseID) REFERENCES Course (CourseID)
        ON UPDATE CASCADE ON DELETE RESTRICT,
    FOREIGN KEY (ExamBoardID) REFERENCES ExamBoard (ExamBoardID)
        ON UPDATE CASCADE ON DELETE RESTRICT
);
"""
(ROOT / "unit7_schema.sql").write_text(schema_sql, encoding="utf-8")

database_path = ROOT / "unit7_normalised.db"
database_path.unlink(missing_ok=True)
connection = sqlite3.connect(database_path)
connection.execute("PRAGMA foreign_keys = ON")
connection.executescript(schema_sql)
connection.executemany(
    "INSERT INTO Student VALUES (:StudentNumber, :StudentName, :ExamScore, :Support, :DateOfBirth)",
    student_rows,
)
connection.executemany(
    "INSERT INTO Teacher VALUES (:TeacherID, :TeacherName)", teacher_rows
)
connection.executemany(
    "INSERT INTO ExamBoard VALUES (:ExamBoardID, :ExamBoardName)", board_rows
)
connection.executemany(
    "INSERT INTO Course VALUES (:CourseID, :CourseName, :TeacherID)", course_rows
)
connection.executemany(
    "INSERT INTO Enrolment VALUES (:StudentNumber, :CourseID, :ExamBoardID)",
    enrolment_rows,
)
connection.commit()

checks: list[tuple[str, int, int]] = []
for table, expected_count in [
    ("Student", 5),
    ("Teacher", 5),
    ("ExamBoard", 5),
    ("Course", 5),
    ("Enrolment", 15),
]:
    actual = connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
    checks.append((f"{table} row count", actual, expected_count))

orphan_queries = {
    "Enrolment rows without a Student": """
        SELECT COUNT(*) FROM Enrolment e
        LEFT JOIN Student s ON s.StudentNumber = e.StudentNumber
        WHERE s.StudentNumber IS NULL
    """,
    "Enrolment rows without a Course": """
        SELECT COUNT(*) FROM Enrolment e
        LEFT JOIN Course c ON c.CourseID = e.CourseID
        WHERE c.CourseID IS NULL
    """,
    "Enrolment rows without an ExamBoard": """
        SELECT COUNT(*) FROM Enrolment e
        LEFT JOIN ExamBoard b ON b.ExamBoardID = e.ExamBoardID
        WHERE b.ExamBoardID IS NULL
    """,
    "Course rows without a Teacher": """
        SELECT COUNT(*) FROM Course c
        LEFT JOIN Teacher t ON t.TeacherID = c.TeacherID
        WHERE t.TeacherID IS NULL
    """,
    "Duplicate student-course pairs": """
        SELECT COUNT(*) FROM (
            SELECT StudentNumber, CourseID, COUNT(*) AS n
            FROM Enrolment GROUP BY StudentNumber, CourseID HAVING n > 1
        )
    """,
}
for label, query in orphan_queries.items():
    checks.append((label, connection.execute(query).fetchone()[0], 0))

rejection_results: list[tuple[str, bool]] = []
for label, sql, params in [
    (
        "Foreign key rejects an unknown student",
        "INSERT INTO Enrolment VALUES (?, ?, ?)",
        (9999, "C01", "B01"),
    ),
    (
        "Composite primary key rejects a duplicate enrolment",
        "INSERT INTO Enrolment VALUES (?, ?, ?)",
        (1001, "C01", "B01"),
    ),
]:
    try:
        connection.execute("SAVEPOINT integrity_test")
        connection.execute(sql, params)
        connection.execute("ROLLBACK TO integrity_test")
        connection.execute("RELEASE integrity_test")
        rejection_results.append((label, False))
    except sqlite3.IntegrityError:
        connection.execute("ROLLBACK TO integrity_test")
        connection.execute("RELEASE integrity_test")
        rejection_results.append((label, True))

joined_count = connection.execute(
    """
    SELECT COUNT(*)
    FROM Enrolment e
    JOIN Student s ON s.StudentNumber = e.StudentNumber
    JOIN Course c ON c.CourseID = e.CourseID
    JOIN Teacher t ON t.TeacherID = c.TeacherID
    JOIN ExamBoard b ON b.ExamBoardID = e.ExamBoardID
    """
).fetchone()[0]
checks.append(("Fully joined enrolment rows", joined_count, 15))

all_passed = all(actual == expected for _, actual, expected in checks) and all(
    passed for _, passed in rejection_results
)

log_lines = [
    "Unit 7 normalisation and relational database integrity tests",
    "Source: DBD_PCOM7E_Table.xlsx",
    "Assumption: ExamScore is student-level because the source supplies one score per student.",
    "",
]
for label, actual, expected_count in checks:
    status = "PASS" if actual == expected_count else "FAIL"
    log_lines.append(f"{status}: {label} = {actual} (expected {expected_count})")
for label, passed in rejection_results:
    log_lines.append(f"{'PASS' if passed else 'FAIL'}: {label}")
log_lines.extend(["", f"Overall result: {'PASS' if all_passed else 'FAIL'}"])
(ROOT / "unit7_integrity_tests.txt").write_text("\n".join(log_lines) + "\n", encoding="utf-8")

notes = """# Unit 7 normalisation notes

## Source assumption

The source supplies one exam score per student while each student has several courses. The score is therefore treated as a student-level attribute. If it instead represents a course assessment, `ExamScore` should move to `Enrolment` and be stored once for each student-course pair.

## Functional dependencies used

- `StudentNumber -> StudentName, ExamScore, Support, DateOfBirth`
- `CourseName -> TeacherName` in the supplied data
- `(StudentNumber, CourseName) -> ExamBoard`

## Normalisation decisions

1. **First normal form:** the repeating course, board and teacher groups were expanded to one row per student-course combination. Values are atomic and the candidate key is `(StudentNumber, CourseName)`.
2. **Second normal form:** attributes depending only on `StudentNumber` moved to `Student`; teacher assignment moved to `Course`; and the student-course relationship retained the board selected for that student.
3. **Third normal form:** the 2NF relations contained no remaining evidenced transitive dependencies. The physical 3NF database uses identifiers for `Teacher`, `Course` and `ExamBoard`, with `Enrolment` linking those entities and enforcing referential integrity.

`Maths` is associated with both AQA and EdExcel in the source. For that reason, exam board is stored on `Enrolment` rather than assumed to depend on course name alone.
"""
(ROOT / "normalisation-notes.md").write_text(notes, encoding="utf-8")

dictionary = """# Unit 7 data dictionary

| Table | Column | Key or constraint | Meaning |
|---|---|---|---|
| Student | StudentNumber | Primary key | Source student identifier |
| Student | StudentName | Required | Student's full name |
| Student | ExamScore | 0–100 | Student-level score supplied in the source |
| Student | Support | Yes/No | Whether support is recorded |
| Student | DateOfBirth | Required | ISO-formatted date of birth |
| Teacher | TeacherID | Primary key | Generated teacher identifier |
| Teacher | TeacherName | Unique | Teacher's display name |
| Course | CourseID | Primary key | Generated course identifier |
| Course | CourseName | Unique | Course name |
| Course | TeacherID | Foreign key | Teacher responsible for the course |
| ExamBoard | ExamBoardID | Primary key | Generated board identifier |
| ExamBoard | ExamBoardName | Unique | Exam board name |
| Enrolment | StudentNumber | Composite primary and foreign key | Enrolled student |
| Enrolment | CourseID | Composite primary and foreign key | Enrolled course |
| Enrolment | ExamBoardID | Foreign key | Board used for that student-course enrolment |
"""
(ROOT / "data-dictionary.md").write_text(dictionary, encoding="utf-8")

connection.close()
if not all_passed:
    raise SystemExit("One or more integrity checks failed.")

print("Built Unit 7 artefacts successfully.")
print(f"1NF rows: {len(one_nf)}")
print(f"3NF rows: {len(student_rows)} students, {len(course_rows)} courses, {len(enrolment_rows)} enrolments")
